"""
Generate ottobiz_pitch_deck.pptx (run from repo root: python scripts/build_ottobiz_pitch_deck.py).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Colors — deep slate + electric coral + sky accent
BG = RGBColor(15, 23, 42)
INK = RGBColor(248, 250, 252)
MUTED = RGBColor(148, 163, 184)
ACCENT = RGBColor(251, 113, 133)
ACCENT2 = RGBColor(56, 189, 248)


def _set_run_font(run, name: str, size: int, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_full_bleed_rect(slide, fill: RGBColor):
    shape = slide.shapes.add_shape(
        1, 0, 0, Inches(13.333), Inches(7.5)
    )  # MSO_SHAPE.RECTANGLE = 1
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()


def title_block(slide, title: str, subtitle: str | None = None, foot: str | None = None):
    add_full_bleed_rect(slide, BG)
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.9), Inches(11.5), Inches(3.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    _set_run_font(p.runs[0], "Segoe UI Light", 44, True, INK)
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.space_before = Pt(18)
        _set_run_font(p2.runs[0], "Segoe UI", 22, False, MUTED)
    if foot:
        p3 = tf.add_paragraph()
        p3.text = foot
        p3.space_before = Pt(28)
        _set_run_font(p3.runs[0], "Segoe UI", 11, False, MUTED)


def content_slide(slide, title: str, bullets: list[str], foot: str | None = None):
    add_full_bleed_rect(slide, BG)
    # Accent bar
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.85), Inches(0.55), Inches(11.5), Inches(0.9))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    _set_run_font(tp.runs[0], "Segoe UI Semibold", 32, True, INK)

    body = slide.shapes.add_textbox(Inches(0.85), Inches(1.45), Inches(11.2), Inches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.space_after = Pt(10)
        para.level = 0
        for run in para.runs:
            _set_run_font(run, "Segoe UI", 17, False, INK)
    if foot:
        fb = slide.shapes.add_textbox(Inches(0.85), Inches(6.55), Inches(11.5), Inches(0.65))
        fp = fb.text_frame.paragraphs[0]
        fp.text = foot
        _set_run_font(fp.runs[0], "Segoe UI", 9, False, MUTED)


def stat_slide(slide, title: str, stats: list[tuple[str, str]], foot: str):
    add_full_bleed_rect(slide, BG)
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT2
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.85), Inches(0.55), Inches(11.5), Inches(0.9))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    _set_run_font(tp.runs[0], "Segoe UI Semibold", 30, True, INK)

    y = 1.35
    for big, small in stats:
        card = slide.shapes.add_shape(1, Inches(0.85), Inches(y), Inches(5.35), Inches(1.25))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        card.line.color.rgb = RGBColor(51, 65, 85)

        tb = slide.shapes.add_textbox(Inches(1.0), Inches(y + 0.12), Inches(5.0), Inches(1.05))
        tf = tb.text_frame
        p0 = tf.paragraphs[0]
        p0.text = big
        _set_run_font(p0.runs[0], "Segoe UI Semibold", 22, True, ACCENT2)
        p1 = tf.add_paragraph()
        p1.text = small
        p1.space_before = Pt(4)
        _set_run_font(p1.runs[0], "Segoe UI", 12, False, MUTED)

        y += 1.42

    fb = slide.shapes.add_textbox(Inches(0.85), Inches(6.5), Inches(11.5), Inches(0.85))
    fp = fb.text_frame.paragraphs[0]
    fp.text = foot
    fp.word_wrap = True
    _set_run_font(fp.runs[0], "Segoe UI", 9, False, MUTED)


def two_col_slide(slide, title: str, left_title: str, left_bullets: list[str], right_title: str, right_bullets: list[str]):
    add_full_bleed_rect(slide, BG)
    tbox = slide.shapes.add_textbox(Inches(0.85), Inches(0.5), Inches(11.5), Inches(0.85))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    _set_run_font(tp.runs[0], "Segoe UI Semibold", 30, True, INK)

    def col(x, y0, col_title: str, items: list[str]):
        ct = slide.shapes.add_textbox(Inches(x), Inches(y0), Inches(5.4), Inches(0.45))
        cp = ct.text_frame.paragraphs[0]
        cp.text = col_title
        _set_run_font(cp.runs[0], "Segoe UI Semibold", 16, True, ACCENT)
        bx = slide.shapes.add_textbox(Inches(x), Inches(y0 + 0.45), Inches(5.4), Inches(5.5))
        tf = bx.text_frame
        for i, line in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.space_after = Pt(8)
            for r in para.runs:
                _set_run_font(r, "Segoe UI", 15, False, INK)

    col(0.85, 1.15, "For shoppers", left_bullets)
    col(6.55, 1.15, "For vendors & ops", right_bullets)


def closing_slide(slide):
    add_full_bleed_rect(slide, BG)
    title_block(
        slide,
        "Ottobiz",
        "AI that sells — and keeps every party in sync.",
        "Thank you.",
    )
    tag = slide.shapes.add_textbox(Inches(0.85), Inches(5.5), Inches(11.5), Inches(0.4))
    tg = tag.text_frame.paragraphs[0]
    tg.text = "Confidential · Draft pitch deck · Figures from public research; verify for due diligence."
    _set_run_font(tg.runs[0], "Segoe UI", 10, False, MUTED)


def main():
    root = Path(__file__).resolve().parents[1]
    out = root / "ottobiz_pitch_deck.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    title_block(
        s1,
        "Ottobiz",
        "The AI-operated commerce layer for your store — chat, pay, deliver, confirm.",
        "Pitch deck · draft",
    )

    # 2 Problem
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s2,
        "Commerce runs in chat — stores still don't",
        [
            "Shoppers discover on Instagram, WhatsApp, DM, and text. The \"store\" is a conversation.",
            "Vendors juggle product answers, payments, disputes, and logistics in fragmented threads — humans burn out, details get lost.",
            "Without a single orchestration layer, conversion leaks at handoff: pricing ambiguity, payment proof, delivery coordination.",
        ],
    )

    # 3 Why people need it
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s3,
        "Why this has to exist",
        [
            "Customers expect instant, persuasive, accurate answers — not \"we'll get back to you\" delays.",
            "SMBs can't hire 24/7 sales + ops for every SKU, receipt, and tracking update.",
             "Trust is won in the thread: proof of payment, transparent logistics, and consistent tone.",
            "The winner is whoever closes the loop: sell → verify pay → create order → ship → upsell — inside one memory.",
        ],
    )

    # 4 Market stats
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    stat_slide(
        s4,
        "Market tailwinds (public signals)",
        [
            (
                "~$10B+ conversational commerce (2024)",
                "Global category projected toward multi‑tens of billions by mid‑2030s (industry sizing).",
            ),
            ("High AI adoption in retail CX", "Majority of retailers piloting or scaling AI for service & conversion (trade press surveys)."),
            ("Chat‑first economies", "In growth markets, social + messaging routinely outpaces \"traditional storefront\" onboarding for MSME sales."),
        ],
        "Sources vary by methodology — treat as directional. Replace with your preferred analyst citations for investor data rooms.",
    )

    # 5 Product
    s5 = prs.slides.add_slide(prs.slide_layouts[6])
    two_col_slide(
        s5,
        "What Ottobiz is",
        "For shoppers",
        [
            "• One persuasive associate that knows your catalog.",
            "• Browse, compare, pay (link or bank transfer), and track.",
            "• Upload receipts & product photos — structured, not chaotic.",
        ],
        "For vendors & ops",
        [
            "• Business console: inventory, analytics, coordination.",
            "• Specialists under the hood: product, payment verify, logistics, complaints.",
            "• Central coordinator + Redis state = durable threads per order/process.",
        ],
    )

    # 6 How it works
    s6 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s6,
        "How it works — agentic, not a basic FAQ bot",
        [
            "Orchestrator reads session memory: active processes, catalog cache, uploads, prior payments.",
            "Handoffs to specialists only when needed — each with slim context, shared state, audit trail.",
            "Payment path supports Paystack + deterministic receipt checks (amount, account, recency).",
            "Logistics hooks: order IDs, tracking, vendor ↔ courier coordination.",
            "Designed for tiered SMB features (inventory upsell, analytics) as you scale.",
        ],
    )

    # 7 Moat / differentiation
    s7 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s7,
        "Why Ottobiz can win",
        [
            "Process‑centric memory: every sale is a thread with IDs — not stateless chat.",
            "Fraud‑aware payment verification + vendor confirmation path — trust sells.",
            "Operates where SMBs already are: messaging‑first, low‑friction onboarding.",
            "Extensible agent mesh: add markets, carriers, POS, ERP without rewriting UX.",
        ],
    )

    # 8 Business model
    s8 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s8,
        "Business model — starter thesis",
        [
            "SaaS tiers per seat / per store (free → gold → platinum feature gating in your code today).",
            "Take rate on payments / affiliate on logistics introductions (optional roadmap).",
            "Add‑ons: analytics packs, priority human‑in‑the‑loop, white‑label storefront widget.",
        ],
        "Tune pricing after pilot CAC/LTV — deck is meant to start conversations.",
    )

    # 9 Roadmap
    s9 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s9,
        "Roadmap (illustrative)",
        [
            "Phase 0–1: Nail shopper joy + vendor reliability (payments, receipts, order creation).",
            "Phase 2: Carrier marketplace density + SLA metrics.",
            "Phase 3: Multichannel inbox (WhatsApp Business API, IG) with shared Ottobiz brain.",
            "Phase 4: Developer API + plugins for POS/inventory leaders.",
        ],
    )

    # 10 Ask
    s10 = prs.slides.add_slide(prs.slide_layouts[6])
    content_slide(
        s10,
        "What we're raising / what we need from you",
        [
            "Capital to scale engineering, trust & safety, and GTM with high‑intent vertical boutiques.",
            "Design partners who live in chat commerce daily — co‑shape receipts, disputes, delivery SLAs.",
            "Strategic intros: payment aggregators, last‑mile fleets, commerce platforms.",
        ],
    )

    s11 = prs.slides.add_slide(prs.slide_layouts[6])
    closing_slide(s11)

    prs.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
